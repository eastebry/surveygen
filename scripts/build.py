#!/usr/bin/env python3
"""Read the translation workbook, compile selected languages, validate, then publish."""
import argparse
import json
import os
from pathlib import Path
import re
import shutil
import subprocess
import sys

from openpyxl import load_workbook
from pypdf import PdfReader

ROOT = Path(os.environ.get('SURVEY_ROOT', Path(__file__).resolve().parents[1]))
LANGUAGES = ('en', 'zh-Hans', 'ar')
UI_KEYS = {'agency', 'survey.title', 'cover.intro', 'cover.instructions.title',
           'cover.instructions', 'cover.start', 'cover.sample', 'instruction.single',
           'instruction.multiple', 'instruction.matrix', 'page.label',
           'footer.continue', 'matrix.statement'}
KEY_PATTERN = re.compile(r'^[a-zA-Z][a-zA-Z0-9_.-]*$')


class SurveyError(ValueError):
    pass


def tex_escape(value):
    """Translate plain text to literal TeX; never interpret workbook text as code."""
    replacements = {'\\': r'\textbackslash{}', '{': r'\{', '}': r'\}',
                    '$': r'\$', '&': r'\&', '#': r'\#', '%': r'\%',
                    '_': r'\_', '^': r'\textasciicircum{}', '~': r'\textasciitilde{}',
                    '\n': r'\newline{}', '\r': '', '\t': ' '}
    return ''.join(replacements.get(c, c) for c in value)


def read_questions(path):
    data = json.loads(path.read_text(encoding='utf-8'))
    questions = data.get('questions', [])
    if not questions:
        raise SurveyError('Question definition must contain at least one question.')
    ids = set()
    for q in questions:
        identifier = q.get('id', '')
        if not KEY_PATTERN.fullmatch(identifier) or identifier in ids:
            raise SurveyError(f'Invalid or duplicate question id: {identifier!r}')
        ids.add(identifier)
        kind = q.get('type')
        if kind not in ('single', 'multiple', 'matrix'):
            raise SurveyError(f'{identifier}: unsupported question type {kind!r}')
        keys = [q.get('prompt', '')]
        if kind == 'matrix':
            if len(q.get('scale', [])) != 5 or not q.get('rows'):
                raise SurveyError(f'{identifier}: matrices require five scale labels and at least one row.')
            keys += q['scale'] + q['rows']
        else:
            options = q.get('options', [])
            if not 2 <= len(options) <= 15:
                raise SurveyError(f'{identifier}: use between 2 and 15 options.')
            keys += options
        if any(not isinstance(k, str) or not KEY_PATTERN.fullmatch(k) for k in keys):
            raise SurveyError(f'{identifier}: invalid translation key.')
    return questions


def required_keys(questions):
    keys = set(UI_KEYS)
    for q in questions:
        keys.add(q['prompt'])
        keys.update(q.get('options', []))
        keys.update(q.get('rows', []))
        keys.update(q.get('scale', []))
    return keys


def read_translations(path, questions, languages):
    wb = load_workbook(path, read_only=True, data_only=False)
    try:
        if 'Translations' not in wb.sheetnames:
            raise SurveyError('Workbook needs a Translations worksheet.')
        rows = iter(wb['Translations'].iter_rows())
        header = [cell.value for cell in next(rows)]
        for name in ('key', 'context', *languages):
            if header.count(name) != 1:
                raise SurveyError(f'Workbook requires exactly one {name!r} column.')
        key_col = header.index('key')
        translations = {lang: {} for lang in languages}
        seen = set()
        for row_number, row in enumerate(rows, 2):
            if all(c.value is None for c in row):
                continue
            key = row[key_col].value
            if not isinstance(key, str) or not KEY_PATTERN.fullmatch(key):
                raise SurveyError(f'Row {row_number}: invalid translation key {key!r}.')
            if key in seen:
                raise SurveyError(f'Row {row_number}: duplicate translation key {key!r}.')
            seen.add(key)
            for lang in languages:
                cell = row[header.index(lang)]
                value = cell.value
                if cell.data_type == 'f':
                    raise SurveyError(f'{key} ({lang}): use plain text, not an Excel formula.')
                if not isinstance(value, str) or not value.strip():
                    raise SurveyError(f'Missing translation: {key} ({lang}), row {row_number}.')
                if any(ord(c) < 32 and c not in '\n\r\t' for c in value):
                    raise SurveyError(f'{key} ({lang}): unsupported control character.')
                if key == 'page.label' and (value.count('{page}') != 1 or value.count('{pages}') != 1):
                    raise SurveyError(f'page.label ({lang}) must contain {{page}} and {{pages}} exactly once.')
                translations[lang][key] = value
        missing = required_keys(questions) - seen
        if missing:
            raise SurveyError('Missing translation keys: ' + ', '.join(sorted(missing)))
        return translations
    finally:
        wb.close()


def generate(root, questions, translations):
    destination = root / 'build/generated'
    destination.mkdir(parents=True, exist_ok=True)
    for lang, strings in translations.items():
        lines = ['% Generated from translations.xlsx. Edit the workbook, not this file.']
        for key, value in strings.items():
            escaped = tex_escape(value)
            if key == 'page.label':
                escaped = escaped.replace(r'\{page\}', r'\thepage{}').replace(r'\{pages\}', r'\getpagerefnumber{survey:last}')
            lines.append(r'\expandafter\newcommand\csname tr@' + key + r'\endcsname{' + escaped + '}')
        (destination / f'strings-{lang}.tex').write_text('\n'.join(lines) + '\n', encoding='utf-8')
    lines = ['% Generated from survey.json. Order is fixed across languages.']
    for number, q in enumerate(questions, 1):
        args = '{' + q['id'] + '}{' + str(number) + '}{' + q['prompt'] + '}'
        last = r'\label{survey:last}' if number == len(questions) else ''
        if q['type'] == 'matrix':
            scale = ''.join('{' + key + '}' for key in q['scale'])
            rows = ''.join(r'\MatrixRow{' + q['id'] + '}{' + str(i) + '}{' + key + '}'
                           for i, key in enumerate(q['rows'], 1))
            lines.append(r'\MatrixQuestion' + args + '{' + scale + '}{' + rows + '}{' + last + '}')
        else:
            options = ''.join(r'\Choice{' + q['id'] + '}{' + str(i) + '}{' + key + '}'
                              for i, key in enumerate(q['options'], 1))
            command = r'\SingleChoiceQuestion' if q['type'] == 'single' else r'\MultipleChoiceQuestion'
            lines.append(command + args + '{' + options + '}{' + last + '}')
    (destination / 'questions.tex').write_text('\n'.join(lines) + '\n', encoding='utf-8')


def read_labels(aux_path):
    return {key: int(page) for key, page in re.findall(
        r'\\newlabel\{([^}]+)\}\{\{[^}]*\}\{(\d+)\}', aux_path.read_text(encoding='utf-8'))}


def validate_pdf(root, language, questions):
    stem = root / 'build' / language / f'survey-{language}'
    log = stem.with_suffix('.log').read_text(encoding='utf-8', errors='replace')
    errors = re.findall(r'^(?:Overfull \\[hv]box.*|Missing character:.*|.*Warning:.*(?:undefined|Rerun|headheight).*)$', log, re.M)
    if errors:
        raise SurveyError(f'{language}: typesetting issues; edit the text or layout.\n' + '\n'.join(errors[:12]))
    pdf = PdfReader(stem.with_suffix('.pdf'))
    page_count = len(pdf.pages) - 1
    if not 2 <= page_count <= 3:
        raise SurveyError(f'{language}: {page_count} question pages; required 2-3. Shorten content or adjust spacing; no text was removed.')
    labels = read_labels(stem.with_suffix('.aux'))
    report = []
    for number, q in enumerate(questions, 1):
        names = [f"q:{q['id']}:start", f"q:{q['id']}:end"]
        count = len(q.get('options', q.get('rows', [])))
        names += [f"a:{q['id']}:{i}" for i in range(1, count + 1)]
        if any(name not in labels for name in names):
            raise SurveyError(f"{language}: missing question or answer marker for {q['id']}.")
        pages = {labels[name] for name in names}
        if len(pages) != 1 or not 1 <= next(iter(pages)) <= page_count:
            raise SurveyError(f"{language}: split or misplaced question {q['id']}.")
        report.append({'number': number, 'id': q['id'], 'type': q['type'],
                       'question_page': next(iter(pages)), 'responses': count})
    if labels.get('survey:last') != page_count:
        raise SurveyError(f'{language}: last-question page reference is incorrect.')
    for page in range(1, page_count + 1):
        if labels.get(f'header:{page}') != page:
            raise SurveyError(f'{language}: missing header on question page {page}.')
        if (f'footer:{page}' in labels) != (page < page_count):
            raise SurveyError(f'{language}: incorrect continuation footer on question page {page}.')
    for i, page in enumerate(pdf.pages):
        if abs(float(page.mediabox.width) - 612) > 1 or abs(float(page.mediabox.height) - 792) > 1:
            raise SurveyError(f'{language}: page {i + 1} is not US Letter.')
        if not page.extract_text().strip():
            raise SurveyError(f'{language}: blank page {i + 1}.')
    fonts = subprocess.check_output(['pdffonts', str(stem.with_suffix('.pdf'))], text=True)
    for line in fonts.splitlines()[2:]:
        # Last columns are: emb sub uni object ID. Names may contain spaces.
        if line.split()[-5] != 'yes':
            raise SurveyError(f'{language}: non-embedded font: {line}')
    return {'language': language, 'question_pages': page_count, 'questions': report}


def compile_language(root, language):
    directory = root / 'build' / language
    directory.mkdir(parents=True, exist_ok=True)
    command = ['xelatex', '-no-shell-escape', '-halt-on-error', '-file-line-error',
               '-interaction=nonstopmode', f'-output-directory={directory}', f'-jobname=survey-{language}',
               r'\def\SurveyLanguage{' + language + r'}\input{survey.tex}']
    for run in range(3):
        result = subprocess.run(command, cwd=root, text=True, stdout=subprocess.PIPE, stderr=subprocess.STDOUT)
        (directory / f'compile-{run + 1}.txt').write_text(result.stdout, encoding='utf-8')
        if result.returncode:
            raise SurveyError(f'{language}: XeLaTeX failed. See build/{language}/compile-{run + 1}.txt\n' + result.stdout[-4500:])


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--language', default=os.environ.get('SURVEY_LANGUAGE', 'all'))
    parser.add_argument('--check-only', action='store_true', help='Validate source data without compiling.')
    args = parser.parse_args()
    languages = LANGUAGES if args.language == 'all' else (args.language,)
    if any(lang not in LANGUAGES for lang in languages):
        raise SurveyError(f'Unsupported language {args.language!r}. Choose all, ' + ', '.join(LANGUAGES))
    questions = read_questions(ROOT / 'survey.json')
    translations = read_translations(ROOT / 'translations.xlsx', questions, languages)
    if args.check_only:
        print(f'Validated {len(questions)} questions and {len(required_keys(questions))} text keys for {", ".join(languages)}.')
        return
    generate(ROOT, questions, translations)
    reports = []
    for lang in languages:
        print(f'Compiling {lang} ...')
        compile_language(ROOT, lang)
        reports.append(validate_pdf(ROOT, lang, questions))
    # Publish only after every requested language passes.
    destination = ROOT / 'output/pdf'
    destination.mkdir(parents=True, exist_ok=True)
    for report in reports:
        lang = report['language']
        shutil.copyfile(ROOT / 'build' / lang / f'survey-{lang}.pdf', destination / f'survey-{lang}.pdf')
        (destination / f'survey-{lang}.validation.json').write_text(json.dumps(report, indent=2) + '\n')
        print(f"Created output/pdf/survey-{lang}.pdf: cover + {report['question_pages']} question pages.")


if __name__ == '__main__':
    try:
        main()
    except (SurveyError, OSError, json.JSONDecodeError) as error:
        print(f'ERROR: {error}', file=sys.stderr)
        sys.exit(1)
