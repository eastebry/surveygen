"""Contract tests, including real workbook edits and XeLaTeX failure cases."""
from collections import Counter
from copy import deepcopy
import json
from pathlib import Path
import shutil
import subprocess
import sys
import tempfile
import unittest
import xml.etree.ElementTree as ET
import zipfile

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'scripts'))
import build

NS = {'s': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}


def edited_workbook(destination, address, value=None, formula=False):
    """Edit a disposable OOXML cell without changing the delivered workbook."""
    with zipfile.ZipFile(ROOT / 'translations.xlsx') as source, zipfile.ZipFile(destination, 'w') as target:
        for info in source.infolist():
            content = source.read(info.filename)
            if info.filename == 'xl/worksheets/sheet1.xml':
                tree = ET.fromstring(content)
                cell = next(c for c in tree.findall('.//s:c', NS) if c.get('r') == address)
                for child in list(cell):
                    cell.remove(child)
                cell.attrib.pop('t', None)
                if formula:
                    ET.SubElement(cell, f"{{{NS['s']}}}f").text = '1+1'
                elif value is not None:
                    cell.set('t', 'inlineStr')
                    inline = ET.SubElement(cell, f"{{{NS['s']}}}is")
                    ET.SubElement(inline, f"{{{NS['s']}}}t").text = value
                content = ET.tostring(tree, encoding='utf-8', xml_declaration=True)
            target.writestr(info, content)


class SurveyTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.work = Path(self.temp.name)
        self.questions = build.read_questions(ROOT / 'survey.json')

    def read_edited(self, address, value=None, formula=False):
        path = self.work / 'edited.xlsx'
        edited_workbook(path, address, value, formula)
        return build.read_translations(path, self.questions, build.LANGUAGES)

    def prepare_compile(self):
        shutil.copyfile(ROOT / 'survey.tex', self.work / 'survey.tex')
        shutil.copytree(ROOT / 'tex', self.work / 'tex')
        return build.read_translations(ROOT / 'translations.xlsx', self.questions, ('en',))

    def test_sample_inventory_and_fixed_order(self):
        self.assertEqual(len(self.questions), 23)
        self.assertEqual(Counter(q['type'] for q in self.questions), {'single': 12, 'multiple': 9, 'matrix': 2})
        self.assertEqual(len(next(q for q in self.questions if q['id'] == 'services')['options']), 15)
        self.assertEqual([len(q['rows']) for q in self.questions if q['type'] == 'matrix'], [3, 3])
        self.assertEqual(self.questions, build.read_questions(ROOT / 'survey.json'))
        self.assertEqual(json.loads((ROOT / 'survey.json').read_text())['shuffle_seed'], 42)

    def test_complete_translations(self):
        strings = build.read_translations(ROOT / 'translations.xlsx', self.questions, build.LANGUAGES)
        self.assertTrue(all(build.required_keys(self.questions) <= set(s) for s in strings.values()))

    def test_missing_translation(self):
        with self.assertRaisesRegex(build.SurveyError, 'Missing translation: agency'):
            self.read_edited('C2')

    def test_duplicate_key(self):
        with self.assertRaisesRegex(build.SurveyError, 'duplicate translation key'):
            self.read_edited('A3', 'agency')

    def test_missing_referenced_key(self):
        with self.assertRaisesRegex(build.SurveyError, 'Missing translation keys: agency'):
            self.read_edited('A2', 'unused.agency')

    def test_missing_language_column(self):
        with self.assertRaisesRegex(build.SurveyError, "'ar' column"):
            self.read_edited('E1', 'other')

    def test_formula_rejected(self):
        with self.assertRaisesRegex(build.SurveyError, 'plain text'):
            self.read_edited('C2', formula=True)

    def test_invalid_page_placeholders(self):
        strings = build.read_translations(ROOT / 'translations.xlsx', self.questions, ('en',))
        # Locate the row by its stable key, not by an assumed row number.
        from openpyxl import load_workbook
        wb = load_workbook(ROOT / 'translations.xlsx', read_only=True)
        row = next(i for i, cells in enumerate(wb['Translations'].iter_rows(values_only=True), 1) if cells[0] == 'page.label')
        wb.close()
        with self.assertRaisesRegex(build.SurveyError, 'exactly once'):
            self.read_edited(f'C{row}', 'Page {page}')

    def test_unknown_language(self):
        result = subprocess.run([sys.executable, str(ROOT / 'scripts/build.py'), '--language', 'unknown'], capture_output=True, text=True)
        self.assertNotEqual(result.returncode, 0)
        self.assertIn('Unsupported language', result.stderr)

    def test_invalid_question_definition(self):
        questions = deepcopy(self.questions)
        questions[0]['type'] = 'unsupported'
        path = self.work / 'survey.json'
        path.write_text(json.dumps({'questions': questions}))
        with self.assertRaisesRegex(build.SurveyError, 'unsupported question type'):
            build.read_questions(path)

    def test_workbook_update_and_literal_tex_symbols_compile(self):
        self.prepare_compile()
        strings = self.read_edited('C8', 'WORKBOOK-UPDATE: 100% & $5_ #1 {A} \\ ~ ^.')
        build.generate(self.work, self.questions, {'en': strings['en']})
        build.compile_language(self.work, 'en')
        build.validate_pdf(self.work, 'en', self.questions)
        extracted = subprocess.check_output(['pdftotext', str(self.work / 'build/en/survey-en.pdf'), '-'], text=True)
        for value in ('WORKBOOK-UPDATE', '100%', '$5_', '#1', '{A}'):
            self.assertIn(value, extracted)

    def test_oversized_question_fails_with_id(self):
        strings = self.prepare_compile()
        strings['en'][self.questions[0]['prompt']] = 'An excessively long question. ' * 1000
        build.generate(self.work, self.questions, strings)
        with self.assertRaisesRegex(build.SurveyError, 'taller than a column'):
            build.compile_language(self.work, 'en')

    def test_page_budget_failure(self):
        strings = self.prepare_compile()
        many = []
        for repeat in range(3):
            for source in self.questions:
                question = deepcopy(source)
                question['id'] += f'-copy{repeat}'
                many.append(question)
        build.generate(self.work, many, strings)
        build.compile_language(self.work, 'en')
        with self.assertRaisesRegex(build.SurveyError, 'question pages; required 2-3'):
            build.validate_pdf(self.work, 'en', many)

    def test_published_languages_and_page_furniture(self):
        for lang in build.LANGUAGES:
            report = build.validate_pdf(ROOT, lang, self.questions)
            self.assertIn(report['question_pages'], (2, 3))
            self.assertEqual([q['id'] for q in report['questions']], [q['id'] for q in self.questions])


if __name__ == '__main__':
    unittest.main(verbosity=2)
